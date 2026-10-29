"""
Exports the client's Registry of Deliverables workbook (TBCA Deliverables Resources
File MP-DEL-2.0, "TBCA_Master_Register_Corrected.xlsx") into the data files read by
DeliverablesRegisterSeeder.

    python3 -I export_register.py <path to TBCA_Master_Register_Corrected.xlsx> [output dir]

Outputs (all in the output directory, default: this folder):

    deliverables.json.gz   the Master Register, 8,318 rows, every column (A to AW)
    crosswalk.csv          the Sprint Crosswalk, 7,589 Sprint codes to deliverable refs
    standards.json         the Standards sheet (tier ladder, size multipliers, page
                           standards by format, CSAT operating rules)
    safeguards.json        the Safeguards sheet (SG-01 to SG-30)
    payment_loops.json     the Payment-Trigger Loops sheet
    integration_specs.json the HeyGen & Infogram sheet
    country_matrix.json    the Country Matrix sheet (54 countries)
    control.json           the Control and Summary sheets

Long texts that repeat across many rows (for example the HeyGen and Infogram layers,
the payment-triggered loop and the acceptance parameters) are stored once in "texts"
and referenced by index, so the file stays small. Nothing is shortened or reworded:
the client's text is carried exactly as written. The per-tab sheets are checked to be
subsets of the Master Register (cell differences are listed in control.json)
and are not exported separately.

The gzip header carries no timestamp, so re-running on the same workbook gives
byte-identical files.
"""

import csv
import gzip
import io
import json
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl

SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else Path('/home/claude/kb/deliverables/TBCA_Master_Register_Corrected.xlsx')
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).parent

# Master Register columns A to AW, in workbook order, with the field name used in the app.
FIELDS = [
    ('ref', 'Ref'),
    ('tab', 'WebApp Tab'),
    ('family', 'Sub-tab / Family'),
    ('name', 'Deliverable'),
    ('format_class', 'Format (11 classes)'),
    ('format_detail', 'Format detail'),
    ('description', 'Description of the deliverable'),
    ('key_issues', 'Key issues to be included'),
    ('extent_basic', 'Extent of detail — BASIC Sprint'),
    ('extent_standard', 'Extent of detail — STANDARD Sprint'),
    ('extent_premium', 'Extent of detail — PREMIUM Sprint'),
    ('pages_basic', 'No. of pages — Basic'),
    ('pages_standard', 'No. of pages — Standard'),
    ('pages_premium', 'No. of pages — Premium'),
    ('page_unit', 'Page unit'),
    ('delivery_basic', 'Delivery — Basic'),
    ('delivery_standard', 'Delivery — Standard'),
    ('delivery_premium', 'Delivery — Premium'),
    ('csat_basic', 'CSAT SLA — Basic (min)'),
    ('csat_standard', 'CSAT SLA — Standard (min)'),
    ('csat_premium', 'CSAT SLA — Premium (min)'),
    ('consulting_standard', 'Consulting standard applied'),
    ('source', 'Source of the deliverable'),
    ('linked_sprint_codes', 'Linked / source Sprint code(s)'),
    ('sprint_type', 'Sprint type (catalogue)'),
    ('audience', 'Who it is for'),
    ('client_problem', 'Client problem solved'),
    ('package_basic', 'Catalogue package — Basic'),
    ('package_standard', 'Catalogue package — Standard'),
    ('package_premium', 'Catalogue package — Premium'),
    ('companion_assets', 'Companion assets'),
    ('client_inputs', 'Client inputs required'),
    ('exclusions', 'Exclusions'),
    ('add_ons', 'Add-ons'),
    ('next_offer', 'Next best offer / upsell'),
    ('catalogue_price', 'Catalogue price B | S | P (ZAR, VAT excl.) and market median'),
    ('catalogue_delivery', 'Catalogue delivery time B | S | P'),
    ('uplift_basic', 'V2 depth uplift — Basic'),
    ('uplift_standard', 'V2 depth uplift — Standard'),
    ('uplift_premium', 'V2 depth uplift — Premium'),
    ('beyond_benchmark', 'Beyond-benchmark value (the wow)'),
    ('ai_depth', 'AI depth — agents, agentic AI, RPA, human'),
    ('heygen_layer', 'HeyGen video layer — Basic | Standard | Premium'),
    ('infogram_layer', 'Infogram visual layer — Basic | Standard | Premium'),
    ('automation_route', 'Automation route (AAI NOW Bible)'),
    ('payment_trigger', 'Payment-triggered delivery loop (no manual start)'),
    ('acceptance', 'Measurable acceptance parameters'),
    ('safeguards', 'Automated safeguards (see Safeguards sheet)'),
    ('v2_status', 'V2 status'),
]

# Columns whose text repeats across many rows: stored once in "texts".
POOLED = [
    'extent_premium', 'consulting_standard', 'audience', 'client_inputs', 'exclusions', 'add_ons',
    'next_offer', 'catalogue_price', 'catalogue_delivery', 'uplift_basic', 'uplift_standard',
    'uplift_premium', 'heygen_layer', 'infogram_layer', 'automation_route', 'payment_trigger', 'acceptance',
]

# Per-tab sheet name -> Master Register "WebApp Tab" value.
TAB_SHEETS = {
    'Services': 'Services', 'Online Shop': 'Online Shop', 'AIIM': 'AIIM', 'Due Diligence': 'Due Diligence',
    'Digital Products': 'Digital Products', 'Subscriptions': 'Subscriptions', 'Retainers': 'Retainers',
    'Marketplace': 'Marketplace', 'Partner Services': 'Partner Services', 'Virtual Deal Room': 'Virtual Deal Room',
    'Technology Solutions': 'Technology Solutions', 'Diagnostics': 'Diagnostics',
    'PA Concept-Bankability': 'Project Advisory — Concept to Bankability',
    'PA Bankability-Turnkey': 'Project Advisory — Bankability to Turnkey',
    'Countries': 'Countries', 'Industries': 'Industries', 'Resources': 'Resources',
    'Tools & Calculators': 'Tools & Calculators', 'Toolkits': 'Toolkits', 'Client Portal': 'Client Portal',
    'KP Thought Leadership': 'KP Thought Leadership', 'Webinars': 'Webinars',
    'Internal Operations': 'Internal Operations (non-client)',
}


def clean(v):
    if v is None:
        return None
    if isinstance(v, float):
        return v
    if isinstance(v, int):
        return v
    s = str(v).strip()
    return s if s != '' else None


def rows_of(ws, first):
    return [tuple(clean(c) for c in r) for r in ws.iter_rows(min_row=first, values_only=True) if any(c is not None for c in r)]


def write_json(name, data):
    (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

    # ---- Master Register -------------------------------------------------
    ws = wb['Master Register']
    header = [clean(c) for c in next(ws.iter_rows(min_row=3, max_row=3, values_only=True))]
    assert header[: len(FIELDS)] == [h for _, h in FIELDS], 'Master Register headers have changed'
    master = [r[: len(FIELDS)] for r in rows_of(ws, 4) if r[0]]
    refs = [r[0] for r in master]
    assert len(refs) == len(set(refs)), 'duplicate deliverable refs'

    texts, text_index = [], {}

    def pool(value):
        if value is None:
            return None
        if value not in text_index:
            text_index[value] = len(texts)
            texts.append(value)
        return text_index[value]

    names = [f for f, _ in FIELDS]
    out_rows = []
    for r in master:
        row = dict(zip(names, r))
        for f in POOLED:
            row[f] = pool(row[f])
        out_rows.append(row)

    # ---- Per-tab sheets are exact subsets of the Master Register ----------
    # Cell differences are reported (the Master Register is the version loaded).
    by_ref = {r[0]: r for r in master}
    tab_check, tab_differences = {}, []
    for sheet, tab in TAB_SHEETS.items():
        sub = [r[: len(FIELDS)] for r in rows_of(wb[sheet], 4) if r[0]]
        expected = [r[0] for r in master if r[1] == tab]
        assert [r[0] for r in sub] == expected, f'{sheet}: refs differ from the Master Register'
        for r in sub:
            for i, (a, b) in enumerate(zip(by_ref[r[0]], r)):
                if a != b:
                    tab_differences.append({'sheet': sheet, 'ref': r[0], 'column': FIELDS[i][1], 'master': a, 'tab_sheet': b})
        tab_check[sheet] = len(sub)

    # ---- Sprint Crosswalk ------------------------------------------------
    cw = [r[:6] for r in rows_of(wb['Sprint Crosswalk'], 4) if r[0]]
    assert len(cw) == len({r[0] for r in cw}), 'duplicate Sprint codes in the crosswalk'
    missing = [r for r in cw if r[4] not in by_ref]
    assert not missing, f'crosswalk refs missing from the register: {missing[:5]}'
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator='\n')
    w.writerow(['sprint_code', 'sprint_name', 'batch', 'population', 'deliverable_ref', 'status'])
    w.writerows(cw)
    (OUT / 'crosswalk.csv').write_text(buf.getvalue(), encoding='utf-8')

    # ---- Standards -------------------------------------------------------
    st = rows_of(wb['Standards'], 1)
    tiers, tiers_v2, sizes, formats, csat_rules = [], [], [], [], []
    section = None
    for r in st[2:]:
        label = r[0]
        if label == 'Parameter' and section is None:
            section = 'tier'
            continue
        if label == 'Issue 2.0 additions to the tier ladder':
            section = 'tier_v2'
            continue
        if label == 'Parameter':
            continue
        if label == 'Size multiplier for page ranges':
            section = 'size'
            continue
        if label == 'Size code':
            continue
        if label == 'Page standards and consulting standard by format':
            section = 'format'
            continue
        if label == 'Format code':
            continue
        if label == 'CSAT SLA operating rule':
            section = 'csat'
            continue
        if section in ('tier', 'tier_v2'):
            (tiers if section == 'tier' else tiers_v2).append(
                {'parameter': label, 'basic': r[1], 'standard': r[2], 'premium': r[3], 'basis': r[4]})
        elif section == 'size':
            sizes.append({'code': label, 'multiplier': r[1], 'meaning': r[2]})
        elif section == 'format':
            pages = [p.strip() for p in str(r[2]).split('|')]
            formats.append({'code': label, 'format': r[1], 'pages_basic': pages[0], 'pages_standard': pages[1],
                            'pages_premium': pages[2], 'unit': r[3], 'consulting_standard': r[4]})
        elif section == 'csat':
            csat_rules.append(label)
    standards = {'title': st[0][0], 'subtitle': st[1][0], 'tiers': tiers, 'tiers_v2': tiers_v2,
                 'sizes': sizes, 'formats': formats, 'csat_rules': csat_rules}

    # Format code for each register row: the Standards format whose detail matches "Format detail".
    detail_to_code = {f['format'].split(' — ', 1)[1].lower(): f['code'] for f in formats}
    unmatched_formats = Counter()
    for row in out_rows:
        code = detail_to_code.get((row['format_detail'] or '').lower())
        row['format_code'] = code
        if not code:
            unmatched_formats[row['format_detail']] += 1

    # ---- Safeguards, loops, integration specs ----------------------------
    sg = rows_of(wb['Safeguards'], 4)
    safeguards = [{'code': r[0], 'name': r[1], 'rule': r[2], 'oss_stage': r[3], 'loop_control': r[4]} for r in sg if r[0]]

    lp = rows_of(wb['Payment-Trigger Loops'], 1)
    loops = {'title': lp[0][0], 'subtitle': lp[1][0], 'loops': [
        {'loop_id': r[0], 'family': r[1], 'task': r[2], 'owning_blueprint': r[3], 'sprints': r[4]} for r in lp[3:] if r[0]]}

    hg = rows_of(wb['HeyGen & Infogram'], 1)
    specs = {'title': hg[0][0], 'subtitle': hg[1][0], 'elements': [
        {'element': r[0], 'heygen': r[1], 'infogram': r[2]} for r in hg[3:] if r[0]]}

    # ---- Country Matrix --------------------------------------------------
    cm = rows_of(wb['Country Matrix'], 1)
    cm_header = list(cm[2])
    countries = []
    for r in cm[3:]:
        if not r[0]:
            continue
        countries.append({
            'iso3': r[0], 'country': r[1], 'au_region': r[2], 'wave': r[3], 'pack_status': r[4],
            'deliverables': {cm_header[i]: r[i] for i in range(5, len(cm_header)) if cm_header[i]},
        })
    country_matrix = {'title': cm[0][0], 'subtitle': cm[1][0], 'deliverables': [h for h in cm_header[5:] if h], 'countries': countries}

    # ---- Control and Summary ---------------------------------------------
    ctrl = rows_of(wb['Control'], 1)
    control = {'title': ctrl[0][0], 'notice': ctrl[1][0],
               'fields': [{'field': r[0], 'value': r[1]} for r in ctrl[3:] if r[0] and r[1]],
               'footer': ctrl[-1][0]}
    sm = rows_of(wb['Summary'], 1)
    sm_header = list(sm[2])
    summary = [dict(zip(sm_header, r)) for r in sm[3:] if r[0]]
    control['summary'] = summary
    control['tab_sheet_differences'] = tab_differences

    # Summary totals must agree with the Master Register.
    total = next(s for s in summary if s['WebApp tab'] == 'Total')
    assert total['Total'] == len(master), 'Summary total differs from the Master Register'
    tab_counts = Counter(r[1] for r in master)
    for s in summary:
        if s['WebApp tab'] != 'Total':
            assert tab_counts[s['WebApp tab']] == s['Total'], f"Summary count differs for {s['WebApp tab']}"

    # ---- Write -----------------------------------------------------------
    doc = {
        'source': SRC.name,
        'issue': next((f['value'] for f in control['fields'] if f['field'] == 'Issue'), None),
        'count': len(out_rows),
        'columns': {f: h for f, h in FIELDS},
        'pooled': POOLED,
        'texts': texts,
        'rows': out_rows,
    }
    raw = json.dumps(doc, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
    with open(OUT / 'deliverables.json.gz', 'wb') as fh:
        with gzip.GzipFile(filename='', mode='wb', fileobj=fh, mtime=0, compresslevel=9) as gz:
            gz.write(raw)

    write_json('standards.json', standards)
    write_json('safeguards.json', safeguards)
    write_json('payment_loops.json', loops)
    write_json('integration_specs.json', specs)
    write_json('country_matrix.json', country_matrix)
    write_json('control.json', control)

    print(f'deliverables: {len(out_rows)} rows, {len(texts)} shared texts, {len(raw):,} bytes before gzip')
    print(f'crosswalk: {len(cw)} Sprint codes; status {dict(Counter(r[5] for r in cw))}')
    print(f'standards: {len(tiers)} tier, {len(tiers_v2)} tier v2, {len(sizes)} sizes, {len(formats)} formats, {len(csat_rules)} CSAT rules')
    print(f'safeguards: {len(safeguards)}; loops: {len(loops["loops"])}; integration elements: {len(specs["elements"])}; countries: {len(countries)}')
    print(f'per-tab sheets checked: {sum(tab_check.values())} rows in {len(tab_check)} sheets; {len(tab_differences)} cell differences')
    for d in tab_differences:
        print(f"  {d['sheet']} {d['ref']} {d['column']}: master {d['master']!r}, tab sheet {d['tab_sheet']!r}")
    if unmatched_formats:
        print(f'format detail without a Standards format code: {dict(unmatched_formats)}')


if __name__ == '__main__':
    main()
